from pathlib import Path
from openpyxl import Workbook, load_workbook


def load_properties(file_path: str | Path) -> dict:
    """Read key=value properties from a .properties file."""
    properties = {}
    with open(file_path, "r", encoding="utf-8") as file:
        for line in file:
            line = line.strip()
            if not line or line.startswith("#") or line.startswith(";"):
                continue
            if "=" in line:
                key, value = line.split("=", 1)
                properties[key.strip()] = value.strip()
    return properties


def resolve_path(base_dir: Path, value: str) -> Path:
    """Resolve a config path relative to the property file directory."""
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = (base_dir / path).resolve()
    return path


def read_excel_from_properties(config_file: str | Path):
    """Read all rows from the Excel file defined in the properties file."""
    config_path = Path(config_file).resolve()
    config_dir = config_path.parent
    config = load_properties(config_path)

    input_path = config.get("input_excel_path")
    sheet_name = config.get("sheet_name")

    if not input_path:
        raise ValueError("input_excel_path is missing in config.properties")

    excel_file = resolve_path(config_dir, input_path)
    if not excel_file.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_file}")

    workbook = load_workbook(excel_file, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    return rows


def write_excel_from_properties(config_file: str | Path, rows):
    """Write rows to the Excel file defined in the properties file."""
    config_path = Path(config_file).resolve()
    config_dir = config_path.parent
    config = load_properties(config_path)

    output_path = config.get("output_excel_path")
    sheet_name = config.get("sheet_name", "Sheet1")

    if not output_path:
        raise ValueError("output_excel_path is missing in config.properties")

    excel_file = resolve_path(config_dir, output_path)
    excel_file.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name

    for row in rows:
        worksheet.append(row)

    workbook.save(excel_file)
    workbook.close()
    print(f"Excel file created successfully: {excel_file}")


if __name__ == "__main__":
    config_file = Path(__file__).with_name("config.properties")

    # Example: read data from Excel
    data = read_excel_from_properties(config_file)
    print("Excel data:")
    for row in data:
        print(row)

    # Example: write new data to Excel
    new_rows = [
        ["Name", "Age", "City"],
        ["Alice", 25, "London"],
        ["Bob", 30, "Paris"],
    ]
    write_excel_from_properties(config_file, new_rows)
